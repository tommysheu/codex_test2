import io
import json
import os
import re
import uuid
from typing import Any, Dict, List, Literal, Optional

import httpx
from docx import Document
from duckduckgo_search import DDGS
from fastapi import FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import HTMLResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from openai import OpenAI
from openpyxl import load_workbook
from pydantic import BaseModel, Field
from pypdf import PdfReader
from pptx import Presentation

app = FastAPI(title="Chatbot with GPT-4o + Web Search", version="1.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

app.mount("/static", StaticFiles(directory="static"), name="static")
templates = Jinja2Templates(directory="templates")

client = OpenAI(api_key=os.getenv("OPENAI_API_KEY"))
upload_store: Dict[str, Dict[str, str]] = {}


class ChatRequest(BaseModel):
    question: str = Field(min_length=1)
    mode: Literal["native", "auto_web", "manual_web"] = "native"
    search_provider: Literal["serpapi", "ddgs"] = "serpapi"
    use_web_search: bool = True
    file_ids: List[str] = Field(default_factory=list)


def detect_need_web_search(question: str) -> bool:
    patterns = [
        r"最新",
        r"news",
        r"today",
        r"current",
        r"即時",
        r"價格",
        r"股價",
        r"天氣",
        r"趨勢",
        r"recent",
    ]
    q = question.lower()
    return any(re.search(p, q) for p in patterns)


def search_with_serpapi(query: str, max_results: int = 5) -> List[Dict[str, str]]:
    api_key = os.getenv("SERPAPI_API_KEY")
    if not api_key:
        raise HTTPException(status_code=400, detail="Missing SERPAPI_API_KEY")

    params = {
        "q": query,
        "engine": "google",
        "api_key": api_key,
        "hl": "zh-tw",
        "num": max_results,
    }
    with httpx.Client(timeout=15) as req:
        res = req.get("https://serpapi.com/search.json", params=params)
        res.raise_for_status()
        data = res.json()

    results: List[Dict[str, str]] = []
    for item in data.get("organic_results", [])[:max_results]:
        results.append(
            {
                "title": item.get("title", ""),
                "url": item.get("link", ""),
                "snippet": item.get("snippet", ""),
            }
        )
    return results


def search_with_ddgs(query: str, max_results: int = 5) -> List[Dict[str, str]]:
    with DDGS() as ddgs:
        rows = ddgs.text(query, max_results=max_results)
        return [
            {
                "title": row.get("title", ""),
                "url": row.get("href", ""),
                "snippet": row.get("body", ""),
            }
            for row in rows
        ]


def extract_text(upload: UploadFile, content: bytes) -> str:
    ext = os.path.splitext(upload.filename or "")[1].lower()

    if ext == ".txt":
        return content.decode("utf-8", errors="ignore")
    if ext == ".pdf":
        reader = PdfReader(io.BytesIO(content))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    if ext in [".docx", ".doc"]:
        doc = Document(io.BytesIO(content))
        return "\n".join(p.text for p in doc.paragraphs)
    if ext in [".ppt", ".pptx"]:
        prs = Presentation(io.BytesIO(content))
        texts: List[str] = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    texts.append(shape.text)
        return "\n".join(texts)
    if ext in [".xls", ".xlsx"]:
        wb = load_workbook(io.BytesIO(content), data_only=True)
        lines: List[str] = []
        for sheet in wb.worksheets:
            lines.append(f"[Sheet] {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                lines.append("\t".join(str(v) if v is not None else "" for v in row))
        return "\n".join(lines)
    if ext in [".png", ".jpg", ".jpeg"]:
        return f"[Image file uploaded: {upload.filename}]"
    raise HTTPException(status_code=400, detail=f"Unsupported file type: {ext}")


def build_messages(payload: ChatRequest, web_results: List[Dict[str, str]]) -> List[Dict[str, Any]]:
    file_context = []
    for file_id in payload.file_ids:
        if file_id in upload_store:
            data = upload_store[file_id]
            file_context.append(f"### File: {data['filename']}\n{data['text'][:8000]}")

    search_context = ""
    if web_results:
        joined = "\n\n".join(
            [f"Title: {r['title']}\nURL: {r['url']}\nSnippet: {r['snippet']}" for r in web_results]
        )
        search_context = f"\n\n[Web Search Results]\n{joined}"

    system_prompt = (
        "你是繁體中文助理。回答請使用 markdown，盡量條列且精簡。"
        "若有提供檔案內容與網頁搜尋結果，請優先引用它們。"
    )

    user_prompt = f"問題：{payload.question}\n\n"
    if file_context:
        user_prompt += "[Uploaded File Context]\n" + "\n\n".join(file_context)
    user_prompt += search_context

    return [
        {"role": "system", "content": system_prompt},
        {"role": "user", "content": user_prompt},
    ]


@app.get("/", response_class=HTMLResponse)
def home(request: Request):
    return templates.TemplateResponse("index.html", {"request": request})


@app.post("/api/upload")
async def upload_file(file: UploadFile = File(...)):
    content = await file.read()
    extracted = extract_text(file, content)
    file_id = str(uuid.uuid4())
    upload_store[file_id] = {"filename": file.filename or "unknown", "text": extracted}
    return {"file_id": file_id, "filename": file.filename, "preview": extracted[:300]}


@app.post("/api/chat/stream")
def chat_stream(payload: ChatRequest):
    web_results: List[Dict[str, str]] = []
    should_search = False

    if payload.mode == "auto_web":
        should_search = detect_need_web_search(payload.question)
    elif payload.mode == "manual_web":
        should_search = payload.use_web_search

    if should_search:
        if payload.search_provider == "serpapi":
            web_results = search_with_serpapi(payload.question)
        else:
            web_results = search_with_ddgs(payload.question)

    messages = build_messages(payload, web_results)

    def stream_gen():
        stream = client.chat.completions.create(
            model="gpt-4o",
            messages=messages,
            stream=True,
            temperature=0.3,
        )
        for chunk in stream:
            delta = chunk.choices[0].delta.content
            if delta:
                yield f"data: {json.dumps({'type': 'token', 'content': delta}, ensure_ascii=False)}\n\n"

        if web_results:
            yield f"data: {json.dumps({'type': 'search_results', 'results': web_results}, ensure_ascii=False)}\n\n"

        yield "data: [DONE]\n\n"

    return StreamingResponse(stream_gen(), media_type="text/event-stream")


@app.get("/health")
def health():
    return {"ok": True}
